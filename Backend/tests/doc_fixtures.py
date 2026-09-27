"""Builds small real documents in memory for upload tests (no binary
fixtures checked in)."""
import io
import zipfile


def make_pdf(pages, encrypt=False) -> bytes:
    """A minimal valid PDF with one text line per page (None = a page with
    no text, like a scanned page). Offsets in the xref table are exact."""
    objects = []

    def add(body: str) -> int:
        objects.append(body)
        return len(objects)

    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    pages_id_placeholder = len(objects) + 1 + 2 * len(pages)
    for text in pages:
        stream = "" if text is None else f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        content = add(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        page_ids.append(add(
            f"<< /Type /Page /Parent {pages_id_placeholder} 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content} 0 R /Resources << /Font << /F1 {font} 0 R >> >> >>"
        ))
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    pages_id = add(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>")
    assert pages_id == pages_id_placeholder
    catalog = add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>")
    encrypt_id = None
    if encrypt:
        encrypt_id = add(
            "<< /Filter /Standard /V 1 /R 2 /P -4 "
            "/O <" + "11" * 32 + "> /U <" + "22" * 32 + "> >>"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    trailer = f"<< /Size {len(objects) + 1} /Root {catalog} 0 R"
    if encrypt_id:
        trailer += f" /Encrypt {encrypt_id} 0 R /ID [<{'ab' * 16}> <{'ab' * 16}>]"
    trailer += " >>"
    out.write(f"trailer\n{trailer}\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs) -> bytes:
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(rows) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_zip_bomb_docx(inflated_bytes: int) -> bytes:
    """A .docx-shaped zip whose single member inflates to `inflated_bytes`."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * inflated_bytes)
    return buf.getvalue()
