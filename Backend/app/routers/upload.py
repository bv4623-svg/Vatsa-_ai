from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from typing import Optional
from datetime import datetime, timezone
import uuid
import io
import os
import logging
from sqlalchemy.orm import Session
from app.auth.dependencies import get_current_user_optional
from app.database import get_db
from app.models.user import User
from app.services.library import register_item, check_quota
from app.services.account import notify_quota_warning
from app.services.storage import get_storage_backend

logger = logging.getLogger("UploadRouter")
router = APIRouter(prefix="/api", tags=["upload"])

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
UPLOAD_STORAGE_ROOT = os.path.join(os.getenv("DATA_DIR") or BACKEND_DIR, "uploads")
# Shared across this app -- app/routers/library/{download,preview}.py and
# app/services/library/deletion.py import this same instance so an upload's
# read/delete path goes through the identical backend (local disk or S3/R2)
# its write path used, instead of each reaching into the local filesystem
# on its own and silently breaking the moment STORAGE_BACKEND=s3 is set.
upload_storage = get_storage_backend(UPLOAD_STORAGE_ROOT)


def sanitize_filename(raw: str) -> str:
    """A client-supplied filename must never become a path component as-is
    -- "../../../etc/passwd" or an absolute path ("C:\\Windows\\...") would
    otherwise let an upload write (and later read back) files outside
    UPLOAD_STORAGE_ROOT. Strips directory separators from both OS
    conventions (this runs on Windows dev machines too) and NUL bytes,
    keeps only the final path segment, and falls back to a safe default
    if nothing usable is left."""
    name = (raw or "").replace("\x00", "")
    name = name.replace("\\", "/").split("/")[-1].strip()
    name = name.lstrip(".") or "file"
    return name[:255]


def thumbnail_key_for(storage_path: str) -> str:
    return f"{storage_path}.thumb.jpg"


# Real file-format signatures, checked against the actual bytes rather than
# trusting the extension a client claims -- a renamed .exe or script can't
# masquerade as one of these just by getting a matching filename. Not
# checked for text files or .svg (both are legitimately arbitrary text,
# with no fixed byte signature to check against).
_MAGIC_BYTES = {
    ".pdf": (b"%PDF-",),
    ".docx": (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"),
    ".xlsx": (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"),
    ".xlsm": (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"),
    ".png": (b"\x89PNG\r\n\x1a\n",),
    ".jpg": (b"\xff\xd8\xff",),
    ".jpeg": (b"\xff\xd8\xff",),
    ".gif": (b"GIF87a", b"GIF89a"),
}

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".log", ".py", ".js", ".jsx", ".ts", ".tsx",
    ".html", ".htm", ".css", ".scss", ".sass", ".xml", ".yaml", ".yml", ".sh", ".bash", ".zsh",
    ".sql", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".kt",
    ".swift", ".dart", ".r", ".m", ".pl", ".lua", ".vue", ".svelte", ".tsv",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm"}

_MIME_BY_EXT = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".gif": "image/gif", ".webp": "image/webp", ".svg": "image/svg+xml",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _check_magic_bytes(raw: bytes, ext: str) -> None:
    signatures = _MAGIC_BYTES.get(ext)
    if not signatures:
        return
    if ext == ".webp":
        return  # handled separately below (RIFF header + WEBP fourcc at offset 8)
    if not any(raw.startswith(sig) for sig in signatures):
        raise HTTPException(400, f"File content doesn't match its {ext} extension")


def _check_webp(raw: bytes) -> None:
    if not (raw.startswith(b"RIFF") and raw[8:12] == b"WEBP"):
        raise HTTPException(400, "File content doesn't match its .webp extension")


def _check_svg(raw: bytes) -> None:
    head = raw[:512].lstrip()
    if not (head.startswith(b"<?xml") or head.startswith(b"<svg")):
        raise HTTPException(400, "File content doesn't look like a valid SVG")


def _make_thumbnail(raw: bytes) -> Optional[bytes]:
    """A small JPEG preview for an uploaded image, or None if Pillow can't
    decode it (corrupt file, or a format Pillow doesn't support -- SVG is
    vector, not decoded here, so it never gets a raster thumbnail)."""
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        img = Image.open(io.BytesIO(raw))
        img.load()
        img.thumbnail((320, 320))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=82)
        return buf.getvalue()
    except Exception:
        logger.warning("Thumbnail generation failed", exc_info=True)
        return None


def extract_text_from_bytes(raw: bytes, filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        try:
            import pdfplumber
        except ImportError:
            raise HTTPException(500, "PDF support not installed. Run: pip install pdfplumber")
        text = ""
        try:
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        except Exception as e:
            logger.error(f"pdfplumber failed on {filename}: {e}")
        return text

    if lower.endswith(".docx"):
        try:
            import docx as _docx
        except ImportError:
            raise HTTPException(500, "DOCX support not installed. Run: pip install python-docx")
        try:
            doc = _docx.Document(io.BytesIO(raw))
            parts = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:
            logger.error(f"python-docx failed on {filename}: {e}")
            return ""

    if lower.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise HTTPException(500, "Excel support not installed. Run: pip install openpyxl")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            parts = []
            for sheet in wb.worksheets:
                parts.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        parts.append(", ".join("" if c is None else str(c) for c in row))
            return "\n".join(parts)
        except Exception as e:
            logger.error(f"openpyxl failed on {filename}: {e}")
            return ""

    raise HTTPException(400, f"Unsupported file type: {filename}")


@router.post("/upload")
async def upload_file(
    file: UploadFile = File(...),
    current_user: Optional[User] = Depends(get_current_user_optional),
    db: Session = Depends(get_db),
):
    raw = await file.read()
    if len(raw) > 25 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 25MB)")
    if len(raw) == 0:
        raise HTTPException(400, "File is empty")

    filename = sanitize_filename(file.filename or "file")
    ext = os.path.splitext(filename)[1].lower()

    is_image = ext in IMAGE_EXTENSIONS
    is_svg = ext == ".svg"
    is_text = ext in TEXT_EXTENSIONS
    is_document = ext in DOCUMENT_EXTENSIONS

    if not (is_image or is_svg or is_text or is_document):
        raise HTTPException(400, f"Unsupported file type: {filename}")

    if ext == ".webp":
        _check_webp(raw)
    elif is_svg:
        _check_svg(raw)
    else:
        _check_magic_bytes(raw, ext)

    # Signed-in users: block before writing, once their Library is full.
    # Anonymous uploads (no account, nothing persisted) can't be measured
    # against a plan, so they're unaffected -- same as before this check.
    if current_user:
        allowed, usage = check_quota(db, current_user, len(raw))
        if not allowed:
            notify_quota_warning(db, current_user, usage["used_bytes"], usage["limit_bytes"], at_limit=True)
            raise HTTPException(status_code=413, detail={
                "error": "storage_limit_reached",
                "used_bytes": usage["used_bytes"],
                "limit_bytes": usage["limit_bytes"],
                "upgrade_url": "/pricing",
            })
        elif usage["at_warning"]:
            notify_quota_warning(db, current_user, usage["used_bytes"], usage["limit_bytes"], at_limit=False)

    text = ""
    if is_document:
        try:
            text = extract_text_from_bytes(raw, filename)
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(400, f"Could not read file: {e}")
    elif is_text:
        text = raw.decode("utf-8", errors="ignore")
    # images/svg: no text extraction -- the chat vision path sends the
    # actual image bytes to the model instead, handled client-side.
    text = (text or "").strip()[:50000]

    file_id = str(uuid.uuid4())
    mime_type = file.content_type or _MIME_BY_EXT.get(ext, "application/octet-stream")

    item = None
    thumbnail_persisted = False
    if current_user:
        storage_path = os.path.join(str(current_user.id), f"{file_id}_{filename}")
        try:
            upload_storage.put(storage_path, raw)
        except Exception:
            logger.exception("Storage write failed for upload %s", file_id)
            raise HTTPException(502, "Could not save the uploaded file. Please try again.")

        if is_image:
            thumb = _make_thumbnail(raw)
            if thumb:
                try:
                    upload_storage.put(thumbnail_key_for(storage_path), thumb)
                    thumbnail_persisted = True
                except Exception:
                    logger.warning("Thumbnail write failed for upload %s", file_id, exc_info=True)

        item = register_item(
            db, current_user.id, "upload", name=filename, size_bytes=len(raw),
            mime=mime_type, source_table="uploads", source_id=file_id,
            storage_path=storage_path,
        )

    return {
        "id": item.id if item else file_id,
        "file_id": file_id,  # back-compat with earlier callers keyed on file_id
        "filename": filename,
        "size": len(raw),
        "size_bytes": len(raw),  # back-compat
        "mime_type": mime_type,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "url": f"/api/library/items/{item.id}/download" if item else None,
        "thumbnail_url": f"/api/library/items/{item.id}/thumbnail" if (item and thumbnail_persisted) else None,
        "chars": len(text),  # back-compat
        "text": text,
    }
