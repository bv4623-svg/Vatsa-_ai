import io

from fastapi import HTTPException

from app.routers.upload.sanitize import _check_magic_bytes
from app.routers.upload.router import logger

try:
    import pdfplumber
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    import docx as _docx
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

try:
    import openpyxl
    XLSX_SUPPORT = True
except ImportError:
    XLSX_SUPPORT = False


def extract_text_from_bytes(raw: bytes, filename: str) -> str:
    _check_magic_bytes(raw, filename)
    lower = filename.lower()
    if lower.endswith(".pdf"):
        if not PDF_SUPPORT:
            raise HTTPException(500, "PDF support not installed. Run: pip install pdfplumber")
        text = ""
        try:
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        except Exception as e:
            logger.error(f"pdfplumber failed on {filename}: {e}")
        return text

    if lower.endswith(".docx"):
        if not DOCX_SUPPORT:
            raise HTTPException(500, "DOCX support not installed. Run: pip install python-docx")
        try:
            doc = _docx.Document(io.BytesIO(raw))
            parts = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:
            logger.error(f"python-docx failed on {filename}: {e}")
            return ""

    if lower.endswith((".xlsx", ".xlsm")):
        if not XLSX_SUPPORT:
            raise HTTPException(500, "Excel support not installed. Run: pip install openpyxl")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            parts = []
            for sheet in wb.worksheets:
                parts.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        parts.append(", ".join("" if c is None else str(c) for c in row))
            return "\n".join(parts)
        except Exception as e:
            logger.error(f"openpyxl failed on {filename}: {e}")
            return ""

    if lower.endswith((".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".html", ".css", ".jsx", ".tsx", ".yaml", ".yml")):
        return raw.decode("utf-8", errors="ignore")

    raise HTTPException(400, f"Unsupported file type: {filename}")
